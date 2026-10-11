"""Aplica las cifras finales de los reportes sin conservar desgloses obsoletos."""
from pathlib import Path
from datetime import datetime
import os
import shutil
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def main():
    target = ROOT / 'Comercial.xlsx'
    backup = ROOT / 'backups' / f'Comercial_final_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
    shutil.copy2(target, backup)
    wb = openpyxl.load_workbook(target)
    july = wb['Julio']
    # Solo filas de pauta: conservar las oportunidades y cotizaciones existentes.
    for row in (2, 3):
        for col in range(3, 9):
            july.cell(row, col).value = None
    july['H2'] = 290857
    for col, region, value in [(15, 'Norte', 107450), (16, 'Centro', 111364), (17, 'Sur', 72043)]:
        july.cell(1, col, f'Meta Ads regional · {region}')
        july.cell(2, col, value)
    notes = wb['Observaciones reportes']
    for row in notes.iter_rows(min_row=2):
        if row[0].value == 'Julio' and 'Inversiones aproximadas' in str(row[1].value):
            row[1].value = ('Reporte final prioritario: Norte 107450, Centro 111364, Sur 72043; '
                            'inversión incorporada 290857 ARS. La fuente presenta estos importes como '
                            'aproximados. Conversaciones y alcance publicitario no informados; '
                            'los ceros técnicos del resumen no acreditan ausencia de resultados.')
    may = wb['Mayo']
    for row in (2, 3):
        for col in range(3, 9):
            may.cell(row, col).value = None
    may['H2'] = 437907.07
    may['H3'] = '(156 conversaciones)'
    may['O1'] = 'Campaña: Meta Ads · campaña 1'
    may['O2'] = 216968.76
    may['O3'] = '(79 conversaciones)'
    may['P1'] = 'Campaña: Meta Ads · Norte y Mixto/Sur'
    may['P2'] = 220938.31
    may['P3'] = '43424 (77 conversaciones)'
    for row in notes.iter_rows(min_row=2):
        if row[0].value == 'Mayo':
            row[1].value = ('Dos campañas distintas de mayo, confirmado por el usuario. '
                            'Meta Ads · campaña 1: 216968,76 ARS y 79 conversaciones; '
                            'Meta Ads · Norte y Mixto/Sur: 220938,31 ARS, 77 conversaciones y '
                            '43424 de alcance bruto. Total: 437907,07 ARS y 156 conversaciones. '
                            'Nombres descriptivos: el Markdown no contiene los nombres oficiales de Meta. '
                            'Alcance de campaña 1 no informado: el alcance agregado disponible es parcial. '
                            'Se toma 79 del encabezado de la primera campaña; su párrafo repite 77.')
    temporary = target.with_name('Comercial.final.tmp.xlsx')
    try:
        wb.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    print(f'Corregido. Respaldo: {backup}')


if __name__ == '__main__':
    main()
