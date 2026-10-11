import pandas as pd
from openpyxl import Workbook
from finance_app.commercial_analytics import load_and_analyze_commercial_data
from finance_app.monthly_reports import available_months, build_monthly_report


def test_august_aggregate_and_social_metrics(tmp_path):
    path = tmp_path / 'commercial.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Agosto'
    ws.append(['AGOSTO', 'Meta Ads sin desglose'])
    ws.append(['Inversion', 135329.58])
    ws.append(['Alcance', '(50 conversaciones)'])
    social = wb.create_sheet('Redes sociales')
    social.append(['Año', 'Mes', 'Visualizaciones'])
    social.append([2026, 'Agosto', 31962])
    wb.save(path)
    result = load_and_analyze_commercial_data(path)
    assert result.kpis.total_inversion_ars == 135329.58
    assert result.kpis.total_conversaciones == 50
    assert result.kpis.total_alcance == 0  # No sustituir alcance por visualizaciones.
    assert result.kpis.cpl_promedio_ars == 135329.58 / 50
    assert available_months(result) == ['Agosto']
    assert result.df_social.iloc[0]['Visualizaciones'] == 31962


def test_two_may_campaigns_replace_legacy_product_breakdown(tmp_path):
    path = tmp_path / 'commercial.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Mayo'
    ws.append(['MAYO', 'TOTAL', 'Campaña: Meta Ads · campaña 1',
               'Campaña: Meta Ads · Norte y Mixto/Sur'])
    ws.append(['Inversion', 437907.07, 216968.76, 220938.31])
    ws.append(['Alcance', '(156 conversaciones)', '(79 conversaciones)', '43424 (77 conversaciones)'])
    wb.save(path)
    result = load_and_analyze_commercial_data(path)
    assert len(result.df_spend) == 2
    assert round(result.kpis.total_inversion_ars, 2) == 437907.07
    assert result.kpis.total_conversaciones == 156
    assert result.kpis.total_alcance == 43424
